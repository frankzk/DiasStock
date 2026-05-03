import argparse
import re
import unicodedata
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import load_workbook

load_dotenv()

from src.config import Store, load_stores
from src.supabase_client import save_snapshot


OUTPUT_RE = re.compile(
    r"^dias_stock_(?P<store>.*?)(?P<date>\d{4}-\d{2}-\d{2})_(?P<time>\d{4})$"
)

COLUMNS = {
    "Producto": ["Producto"],
    "SKU": ["SKU"],
    "Stock Actual": ["Stock Actual"],
    "Ventas 7d": ["Ventas 7d"],
    "Ventas/dia (avg)": ["Ventas/dia (avg)", "Ventas/día (avg)"],
    "Dias de Stock": ["Dias de Stock", "Días de Stock"],
    "Estado de Inventario": ["Estado de Inventario", "Estado"],
    "Estado de Venta": ["Estado de Venta"],
    "Alerta": ["Alerta", "Accion", "Acción"],
    "Imagen": ["Imagen"],
}


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Importa Excels historicos de outputs/ a Supabase."
    )
    parser.add_argument("--outputs-dir", default="outputs")
    parser.add_argument("--store", help="Filtra por key de tienda, ej: CR, HN, KE, KA")
    parser.add_argument("--date", help="Filtra por fecha YYYY-MM-DD")
    parser.add_argument("--all-runs", action="store_true", help="Procesa todos los Excels, no solo el ultimo por dia/tienda")
    parser.add_argument("--dry-run", action="store_true", help="Solo muestra lo que importaria, sin subir a Supabase")
    args = parser.parse_args()

    stores = load_stores()
    store_by_slug = {_slug(store.name): store for store in stores}
    files = _discover_files(Path(args.outputs_dir), store_by_slug)

    if args.store:
        key = args.store.upper()
        files = [item for item in files if item["store"].key == key]
    if args.date:
        files = [item for item in files if item["date"] == args.date]
    if not args.all_runs:
        files = _latest_file_per_store_date(files)

    if not files:
        print("No se encontraron Excels para importar.")
        return

    print(f"Excels seleccionados: {len(files)}")
    total_rows = 0

    for item in files:
        results = _read_detail_sheet(item["path"])
        total_rows += len(results)
        store = item["store"]
        print(
            f"  {item['date']} {store.key:<3} {store.name:<24} "
            f"{len(results):>4} filas  {item['path'].name}"
        )

        if not args.dry_run:
            save_snapshot(
                results,
                store_key=store.key,
                store_name=store.name,
                run_date=item["date"],
            )

    action = "validaria" if args.dry_run else "importo"
    print(f"\nListo: se {action} {total_rows} filas en {len(files)} snapshot(s).")


def _discover_files(outputs_dir: Path, store_by_slug: dict[str, Store]) -> list[dict]:
    files = []
    for path in outputs_dir.glob("dias_stock_*.xlsx"):
        match = OUTPUT_RE.match(path.stem)
        if not match:
            continue

        raw_store = match.group("store").strip("_")
        slug = _slug(raw_store)
        store = store_by_slug.get(slug)
        if not store:
            print(f"  Omitido sin tienda configurada: {path.name}")
            continue

        files.append(
            {
                "path": path,
                "store": store,
                "date": match.group("date"),
                "time": match.group("time"),
            }
        )

    return sorted(files, key=lambda item: (item["date"], item["store"].key, item["time"]))


def _latest_file_per_store_date(files: list[dict]) -> list[dict]:
    latest = {}
    for item in files:
        key = (item["store"].key, item["date"])
        if key not in latest or item["time"] > latest[key]["time"]:
            latest[key] = item
    return sorted(latest.values(), key=lambda item: (item["date"], item["store"].key))


def _read_detail_sheet(path: Path) -> list[dict]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb["Detalle"] if "Detalle" in wb.sheetnames else wb.active
    header_row = _find_header_row(ws)
    headers = {
        str(cell.value).strip(): index
        for index, cell in enumerate(ws[header_row], start=1)
        if cell.value
    }

    required = ["Producto", "SKU", "Stock Actual", "Ventas 7d", "Ventas/dia (avg)", "Dias de Stock", "Estado de Inventario"]
    missing = [name for name in required if _column(headers, name) is None]
    if missing:
        raise ValueError(f"{path.name}: faltan columnas {', '.join(missing)}")

    results = []
    for row_idx in range(header_row + 1, ws.max_row + 1):
        product = _cell(ws, row_idx, headers, "Producto")
        sku = _cell(ws, row_idx, headers, "SKU")
        if not product and not sku:
            continue

        results.append(
            {
                "Producto": product,
                "Imagen": _cell(ws, row_idx, headers, "Imagen"),
                "SKU": sku,
                "Stock Actual": _number(_cell(ws, row_idx, headers, "Stock Actual"), 0),
                "Ventas 7d": _number(_cell(ws, row_idx, headers, "Ventas 7d"), 0),
                "Ventas/dia (avg)": _number(_cell(ws, row_idx, headers, "Ventas/dia (avg)"), 0),
                "Dias de Stock": _optional_number(_cell(ws, row_idx, headers, "Dias de Stock")),
                "Estado de Inventario": _cell(ws, row_idx, headers, "Estado de Inventario"),
                "Estado de Venta": _cell(ws, row_idx, headers, "Estado de Venta"),
                "Alerta": _cell(ws, row_idx, headers, "Alerta"),
            }
        )

    _fill_missing_sale_status(results)
    return results


def _find_header_row(ws) -> int:
    for row_idx in range(1, min(ws.max_row, 8) + 1):
        values = {str(cell.value).strip() for cell in ws[row_idx] if cell.value}
        if {"Producto", "SKU", "Stock Actual"}.issubset(values):
            return row_idx
    raise ValueError(f"{ws.title}: no se encontro la fila de encabezados")


def _cell(ws, row_idx: int, headers: dict[str, int], name: str) -> str:
    col_idx = _column(headers, name)
    if not col_idx:
        return ""
    value = ws.cell(row=row_idx, column=col_idx).value
    return "" if value is None else str(value).strip()


def _column(headers: dict[str, int], name: str) -> int | None:
    for candidate in COLUMNS.get(name, [name]):
        if candidate in headers:
            return headers[candidate]
    return None


def _fill_missing_sale_status(results: list[dict]) -> None:
    max_daily = max((row["Ventas/dia (avg)"] for row in results if row["Ventas/dia (avg)"] > 0), default=1)
    for row in results:
        if row["Estado de Venta"]:
            continue
        daily_avg = row["Ventas/dia (avg)"]
        if daily_avg == 0:
            row["Estado de Venta"] = "SIN VENTAS"
            continue
        ratio = daily_avg / max_daily if max_daily > 0 else 0
        if ratio >= 0.40:
            row["Estado de Venta"] = "ALTA"
        elif ratio >= 0.10:
            row["Estado de Venta"] = "MEDIA"
        else:
            row["Estado de Venta"] = "BAJA"


def _number(value: str, default: int | float) -> int | float:
    parsed = _optional_number(value)
    return default if parsed is None else parsed


def _optional_number(value: str) -> int | float | None:
    if value in ("", "-", "None", "Sin ventas"):
        return None
    try:
        number = float(str(value).replace(",", "."))
    except ValueError:
        return None
    return int(number) if number.is_integer() else number


def _slug(value: str) -> str:
    value = unicodedata.normalize("NFC", value.strip().lower())
    return re.sub(r"\s+", "_", value)


if __name__ == "__main__":
    main()
