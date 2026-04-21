import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


STATUS_COLORS = {
    "CRITICO":    "FFCCCC",  # rojo claro
    "BAJO":       "FFE5CC",  # naranja claro
    "OK":         "CCFFCC",  # verde claro
    "ALTO":       "CCE5FF",  # azul claro
    "Sin ventas": "F0F0F0",  # gris
}

HEADER_FILL = PatternFill("solid", fgColor="1F2D3D")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def export_to_excel(results: list[dict], output_dir: str = "outputs") -> str:
    os.makedirs(output_dir, exist_ok=True)
    date_str = datetime.now().strftime("%Y-%m-%d_%H%M")
    filename = os.path.join(output_dir, f"dias_stock_{date_str}.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "Días de Stock"

    columns = list(results[0].keys()) if results else []

    # Headers
    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    ws.row_dimensions[1].height = 22

    # Data rows
    for row_idx, row_data in enumerate(results, start=2):
        status = row_data.get("Estado", "")
        fill_color = STATUS_COLORS.get(status, "FFFFFF")
        row_fill = PatternFill("solid", fgColor=fill_color)

        for col_idx, col_name in enumerate(columns, start=1):
            value = row_data[col_name]
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = row_fill
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center")

            if col_name == "Estado":
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # Auto-fit column widths
    for col_idx, col_name in enumerate(columns, start=1):
        max_len = len(col_name)
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 60)

    # Freeze header row
    ws.freeze_panes = "A2"

    wb.save(filename)
    print(f"  Excel guardado: {filename}")
    return filename
